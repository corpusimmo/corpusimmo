"""
LES APERÇUS DES CLASSEURS, TIRÉS DES FICHIERS EUX-MÊMES.

Une capture par onglet de chaque matrice de `public/outils/matrices/`, écrite
dans `public/outils/apercus/<outil>-<n>.jpg`. La première sert d'illustration à
la fiche de l'outil dans la bibliothèque : elle doit donc montrer le calculateur,
pas le mode d'emploi.

LE RENDU PASSE PAR NUMBERS, ET C'EST UN CHOIX DE FIDÉLITÉ. Aucun moteur
installé ici ne rend un classeur comme un tableur : ni LibreOffice, ni un
export HTML maison. Numbers ouvre le `.xlsx`, recalcule, applique les styles
et exporte chaque feuille sur une page PDF ; PDFKit (`pdf-en-png.swift`) la
rasterise ; PIL rogne les marges et normalise la largeur.

LA COPIE PHOTOGRAPHIÉE EST FIGÉE SUR LES VALEURS D'EXCEL. Numbers recalcule à
l'ouverture, et il ne recalcule pas comme Excel : il reporte le format d'une
cellule référencée dans une concaténation (« 8 % du prix » devenait
« 800,0 % % du prix ») et rend 0 sur des contrôles qu'Excel évalue à « OK ». Un
aperçu qui montre « À CORRIGER » sur l'exemple livré dessert le classeur qu'il
présente. Chaque formule de la copie est donc remplacée par la valeur
qu'Excel a enregistrée dans le fichier : Numbers ne calcule plus rien, il met
en page. Le fichier téléchargeable, lui, garde toutes ses formules.

    python3 scripts/outils/apercus.py
"""

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from PIL import Image, ImageChops

RACINE = Path(__file__).resolve().parents[2]
MATRICES = RACINE / "public/outils/matrices"
SORTIE = RACINE / "public/outils/apercus"
SWIFT = RACINE / "scripts/outils/pdf-en-png.swift"

# Classeur → identifiant d'outil. L'ordre des onglets est celui du fichier.
CLASSEURS = {
    "calculateur-rentabilite-locative.xlsx": "rentabilite-locative",
    "tableau-amortissement-et-comparateur-de-prets.xlsx": "pret-amortissement",
    "capacite-emprunt-et-bilan-patrimonial.xlsx": "capacite-emprunt",
    "arbitrage-fiscal-nu-lmnp-sci-is.xlsx": "arbitrage-fiscal",
    "chiffrage-de-travaux-par-lot.xlsx": "chiffrage-travaux",
    "rent-roll-et-wault.xlsx": "wault",
    "avis-de-valeur-par-comparaison.xlsx": "avis-de-valeur",
    "net-vendeur-honoraires-et-qualification.xlsx": "net-vendeur",
    "dcf-valorisation.xlsx": "dcf",
}

# Le mode d'emploi ne se photographie pas : c'est du texte, que la fiche de
# l'outil dit déjà, et une page de prose ne prouve pas qu'un classeur calcule.
ONGLETS_EXCLUS = {"mode d'emploi"}

LARGEUR = 1400
MARGE = 36


def copie_pour_rendu(source: Path, dossier: Path) -> tuple[Path, list[str]]:
    """Copie le classeur, chaque formule remplacée par sa valeur calculée par Excel."""
    wb = openpyxl.load_workbook(source)
    valeurs = openpyxl.load_workbook(source, data_only=True)
    for ws in wb.worksheets:
        figee = valeurs[ws.title]
        for ligne in ws.iter_rows():
            for cellule in ligne:
                v = cellule.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                calculee = figee[cellule.coordinate].value
                cellule.value = calculee
                # Un texte calculé ne garde pas un format numérique : Numbers
                # l'appliquerait encore au rendu.
                if isinstance(calculee, str):
                    cellule.number_format = "General"
    cible = dossier / "rendu.xlsx"
    wb.save(cible)
    return cible, [ws.title for ws in wb.worksheets]


def exporter_pdf(xlsx: Path, pdf: Path) -> None:
    script = f'''
    set src to POSIX file "{xlsx}"
    set dst to POSIX file "{pdf}"
    tell application "Numbers"
      set d to open src
      delay 2
      export d to dst as PDF
      close d saving no
    end tell'''
    subprocess.run(["osascript", "-e", script], check=True, capture_output=True)


def rogner(image: Image.Image) -> Image.Image:
    fond = Image.new(image.mode, image.size, (255, 255, 255))
    boite = ImageChops.difference(image, fond).getbbox()
    if not boite:
        return image
    g, h, d, b = boite
    return image.crop(
        (max(0, g - MARGE), max(0, h - MARGE), min(image.width, d + MARGE), min(image.height, b + MARGE))
    )


def main() -> None:
    releve = {}
    for fichier, outil in CLASSEURS.items():
        source = MATRICES / fichier
        if not source.exists():
            print(f"  absent : {fichier}", file=sys.stderr)
            continue
        with tempfile.TemporaryDirectory() as tmp:
            dossier = Path(tmp)
            xlsx, onglets = copie_pour_rendu(source, dossier)
            pdf = dossier / "rendu.pdf"
            exporter_pdf(xlsx, pdf)
            pages = dossier / "pages"
            pages.mkdir()
            subprocess.run(["swift", str(SWIFT), str(pdf), str(pages), "2"], check=True, capture_output=True)
            images = sorted(pages.glob("page-*.png"))

            if len(images) != len(onglets):
                print(
                    f"  {fichier} : {len(images)} pages pour {len(onglets)} onglets, "
                    "correspondance onglet/page incertaine",
                    file=sys.stderr,
                )

            for ancien in SORTIE.glob(f"{outil}-*.jpg"):
                ancien.unlink()

            retenus = []
            for onglet, png in zip(onglets, images):
                if onglet.strip().lower() in ONGLETS_EXCLUS:
                    continue
                image = rogner(Image.open(png).convert("RGB"))
                hauteur = round(image.height * LARGEUR / image.width)
                image = image.resize((LARGEUR, hauteur), Image.LANCZOS)
                rang = len(retenus) + 1
                image.save(SORTIE / f"{outil}-{rang}.jpg", "JPEG", quality=86, optimize=True, progressive=True)
                retenus.append([onglet, LARGEUR, hauteur])
            releve[outil] = retenus
            print(f"  {outil} : {len(retenus)} captures", file=sys.stderr)

    print(json.dumps(releve, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
